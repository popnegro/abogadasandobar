from docx import Document

from docx.shared import Pt, Inches

from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx.enum.table import WD_TABLE_ALIGNMENT



doc = Document()

sec = doc.sections[0]

sec.top_margin = Inches(0.65)

sec.bottom_margin = Inches(0.65)

sec.left_margin = Inches(0.75)

sec.right_margin = Inches(0.75)



styles = doc.styles

styles["Normal"].font.name = "Aptos"

styles["Normal"].font.size = Pt(10)



p = doc.add_paragraph()

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = p.add_run("INFORME TÉCNICO DE ESTADO Y PLAN DE MEJORA")

r.bold = True

r.font.size = Pt(19)



p = doc.add_paragraph()

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

r = p.add_run("Emilia Sandobar · Abogada Penalista en Mendoza")

r.bold = True

r.font.size = Pt(13)



p = doc.add_paragraph()

p.alignment = WD_ALIGN_PARAGRAPH.CENTER

p.add_run("Baseline inicial · 18 de septiembre de 2026").italic = True



doc.add_heading("1. Objetivo", level=1)

doc.add_paragraph(

    "Establecer un punto de referencia técnico y de posicionamiento para medir durante los próximos "

    "30 días la evolución de la presencia digital de Emilia Sandobar. Esta etapa se concentra en "

    "SEO, web, Google Business Profile (GBP) y GEO."

)



doc.add_heading("2. Estado actual", level=1)

doc.add_paragraph(

    "La infraestructura técnica del sitio se encuentra consolidada, con un baseline previo de 100/100 "

    "en las métricas principales de PageSpeed y 3/3 en navegación con agentes. La prioridad actual no "

    "es rediseñar la web, sino consolidar las páginas estratégicas, mejorar la calidad de medición y "

    "generar señales SEO, locales y GEO comparables."

)



doc.add_heading("3. Baseline cuantitativo — 18/09/2026", level=1)

table = doc.add_table(rows=1, cols=4)

table.style = "Table Grid"

table.alignment = WD_TABLE_ALIGNMENT.CENTER

for i,t in enumerate(["Fuente","Indicador","Baseline","Observación"]):

    table.rows[0].cells[i].text=t



baseline = [

("Search Console","Clics orgánicos · últimos 28 días","0","Datos asentados hasta 16/09"),

("Search Console","Impresiones · últimos 28 días","21","Datos asentados hasta 16/09"),

("Search Console","CTR","0%","Sin clics registrados en la ventana"),

("Search Console","Posición media","10,61","Promedio general"),

("Search Console","Consulta 'abogada penalista'","Pos. 5–6","Señal inicial sobre /abogada-penalista-mendoza"),

("Search Console","/abogada-penalista-mendoza","15 impresiones","URL prioritaria"),

("Search Console","/servicios-abogacia-mendoza","10 impresiones","Posición media aproximada 7,8"),

("GBP","Impresiones recientes","1","Dato disponible del 15/09"),

("GBP","Clics web / llamadas / indicaciones","0 / 0 / 0","En el dato reciente disponible"),

("GBP","Reseñas","1 · 5,0","Dato actual disponible"),

("GEO","Navegación con agentes","3/3","Baseline previo del proyecto"),

]

for row in baseline:

    cells=table.add_row().cells

    for i,t in enumerate(row): cells[i].text=t



doc.add_paragraph(

    "Nota metodológica: Search Console presenta desfase de datos. Por ello, el documento se fecha "

    "el 18/09/2026, mientras que los datos SEO asentados utilizados llegan hasta el 16/09/2026."

).italic = True



doc.add_heading("4. Punto 1 — Auditoría de GA4", level=1)

doc.add_paragraph(

    "La primera acción es validar la medición antes de utilizar conversiones como KPI comercial. "

    "La propiedad presenta señales de tráfico de desarrollo/pruebas, incluyendo actividad desde "

    "127.0.0.1. También aparecen conversiones anormalmente perfectas en algunas rutas de contacto. "

    "Por lo tanto, esas conversiones no deben interpretarse todavía como consultas comerciales."

)

for x in [

"Separar tráfico de producción de pruebas/desarrollo.",

"Identificar exactamente qué eventos están marcados como conversiones.",

"Conservar como conversiones principales únicamente acciones verificables: WhatsApp, teléfono y formulario.",

"Registrar el nuevo baseline GA4 una vez corregida la medición."

]:

    doc.add_paragraph(x, style="List Bullet")



doc.add_heading("5. Punto 2 — Consolidación de páginas prioritarias", level=1)

table = doc.add_table(rows=1, cols=5)

table.style="Table Grid"

for i,t in enumerate(["URL","Indexación","Impresiones","Posición","Acciones"]):

    table.rows[0].cells[i].text=t

for row in [

("/abogada-penalista-mendoza","PASS","15","Señal 5–6 para 'abogada penalista'","H1, contenido, enlaces internos, CTA"),

("/servicios-abogacia-mendoza","PASS","10","≈ 7,8","H1, contenido, enlaces internos, CTA"),

]:

    cells=table.add_row().cells

    for i,t in enumerate(row): cells[i].text=t



doc.add_paragraph(

    "La consolidación será incremental: no se propone un rediseño. El objetivo es mejorar la "

    "relevancia textual, la arquitectura interna y el recorrido hacia contacto manteniendo el sistema visual."

)



doc.add_heading("6. Punto 3 — Google Business Profile", level=1)

doc.add_paragraph(

    "El perfil 'Abogada Penalista Emilia Sandobar' ya está conectado y operativo. El volumen actual "

    "todavía es reducido, por lo que el objetivo de los próximos 30 días es construir histórico comparable."

)

for x in [

"Publicar contenido periódico relacionado con los servicios prioritarios.",

"Reforzar servicios y descripción manteniendo información estable y verificable.",

"Consolidar las señales Emilia Sandobar + abogada penalista + Mendoza.",

"Medir Search, Maps, clics al sitio, llamadas, indicaciones y reseñas.",

"No utilizar datos variables como eje de posicionamiento."

]:

    doc.add_paragraph(x, style="List Bullet")



doc.add_heading("7. Punto 4 — Benchmark GEO", level=1)

doc.add_paragraph(

    "El baseline GEO es 3/3 en navegación con agentes. En 30 días se repetirá el benchmark con "

    "las mismas consultas y criterios para que la comparación sea cuantitativa."

)

for x in [

"Consulta principal: 'abogada penalista en Mendoza'.",

"Registrar si Emilia aparece.",

"Registrar si aparece el sitio oficial.",

"Registrar si la especialidad y ubicación se describen correctamente.",

"Registrar URLs citadas y fecha de ejecución."

]:

    doc.add_paragraph(x, style="List Bullet")



doc.add_heading("8. Plan de ejecución — 30 días", level=1)

table = doc.add_table(rows=1, cols=3)

table.style="Table Grid"

for i,t in enumerate(["Período","Acciones","Resultado esperado"]):

    table.rows[0].cells[i].text=t

for row in [

("Días 1–7","Auditar/corregir GA4 y congelar baseline","Datos de conversión confiables"),

("Días 8–15","Consolidar las dos páginas prioritarias","Mayor relevancia y mejor recorrido"),

("Días 16–23","Activar/optimizar GBP","Más señales locales medibles"),

("Días 24–30","Repetir benchmark GEO y comparar","Evolución contra baseline"),

]:

    cells=table.add_row().cells

    for i,t in enumerate(row): cells[i].text=t



doc.add_heading("9. Indicadores al día 30", level=1)

for x in [

"Search Console: impresiones, clics, CTR y posición media.",

"Consulta 'abogada penalista': impresiones, posición y CTR.",

"/abogada-penalista-mendoza: impresiones, posición y clics.",

"/servicios-abogacia-mendoza: impresiones, posición y clics.",

"GA4: usuarios, engagement y conversiones válidas de producción.",

"GBP: impresiones Search/Maps, clics web, llamadas, indicaciones y reseñas.",

"GEO: resultado del benchmark repetido."

]:

    doc.add_paragraph(x, style="List Bullet")



doc.add_heading("10. Criterio de éxito", level=1)

doc.add_paragraph(

    "El objetivo de esta etapa no es prometer una cantidad determinada de consultas. Es construir "

    "un sistema de medición confiable y mejorar las señales de relevancia de las dos páginas prioritarias, "

    "la presencia local en Google y la recuperación de la entidad Emilia Sandobar en búsquedas GEO."

)

doc.add_paragraph(

    "Al finalizar los 30 días se realizará una comparación contra este baseline para decidir qué acciones "

    "mantener, ampliar o corregir."

)



footer = doc.sections[0].footer.paragraphs[0]

footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

footer.add_run("Emilia Sandobar · Informe técnico · 18/09/2026")



doc.save("Informe_Tecnico_Emilia_Sandobar.docx")

